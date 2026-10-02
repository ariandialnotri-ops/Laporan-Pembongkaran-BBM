#!/usr/bin/env python3
"""
Ubah template Excel laporan (Berita Acara Pembongkaran BBM & Catatan
Persediaan BBM) menjadi data tata letak untuk tampilan HTML yang identik
(dipakai untuk PDF/JPG), dan salin template + logonya ke public/templates/.

    python3 scripts/build-report-templates.py path/ke/template.xlsx

Excel tetap diisi langsung dari file template aslinya (src/lib/report/xlsx.ts),
sehingga font, logo, garis tabel, dan format angka tidak berubah.
"""
import colorsys
import json
import os
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.styles.colors import COLOR_INDEX

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_JSON = os.path.join(ROOT, 'src', 'data', 'report-templates.json')
OUT_PUBLIC = os.path.join(ROOT, 'public', 'templates')

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

# Sheet -> kunci & area cetak (kolom/baris pertama..terakhir).
SHEETS = {
    'BERITA ACARA PEMBONGKARAN': ('ba', 'B1:O69'),
    'Catatan Persediaan BBM,': ('persediaan', 'A1:T37'),
}

BORDER_CSS = {
    'hair': '1px dotted',
    'dotted': '1px dotted',
    'dashDot': '1px dashed',
    'dashDotDot': '1px dashed',
    'dashed': '1px dashed',
    'thin': '1px solid',
    'mediumDashed': '2px dashed',
    'mediumDashDot': '2px dashed',
    'mediumDashDotDot': '2px dashed',
    'slantDashDot': '2px dashed',
    'medium': '2px solid',
    'thick': '3px solid',
    'double': '3px double',
}


def theme_palette(xlsx):
    with zipfile.ZipFile(xlsx) as z:
        root = ET.fromstring(z.read('xl/theme/theme1.xml'))
    scheme = root.find('.//a:clrScheme', NS)
    vals = {}
    for el in scheme:
        tag = el.tag.split('}')[1]
        child = el[0]
        vals[tag] = child.get('lastClr') or child.get('val')
    # Urutan indeks tema Excel: lt1, dk1, lt2, dk2, accent1..6, hlink, folHlink.
    order = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    return [vals.get(k, '000000') for k in order]


def apply_tint(hex6, tint):
    if not tint:
        return hex6
    r, g, b = (int(hex6[i : i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))


def color_of(c, palette, default=None):
    if c is None:
        return default
    try:
        if c.type == 'rgb' and c.rgb and isinstance(c.rgb, str):
            return '#' + c.rgb[-6:]
        if c.type == 'theme' and c.theme is not None:
            return '#' + apply_tint(palette[c.theme], c.tint or 0)
        if c.type == 'indexed' and c.indexed is not None:
            if c.indexed >= 64:
                return default
            return '#' + COLOR_INDEX[c.indexed][-6:]
    except Exception:
        return default
    return default


def raw_cols(xlsx, sheet_path):
    """Lebar kolom langsung dari <cols> (openpyxl menggabungkan rentang)."""
    with zipfile.ZipFile(xlsx) as z:
        root = ET.fromstring(z.read(sheet_path))
    fmt = root.find('m:sheetFormatPr', NS)
    default_w = float(fmt.get('defaultColWidth') or fmt.get('baseColWidth') or 8) if fmt is not None else 8.43
    if fmt is not None and fmt.get('defaultColWidth') is None:
        default_w = 8.43 if not fmt.get('baseColWidth') else float(fmt.get('baseColWidth')) + 0.43
    widths, hidden = {}, set()
    for col in root.findall('m:cols/m:col', NS):
        for i in range(int(col.get('min')), int(col.get('max')) + 1):
            widths[i] = float(col.get('width', default_w))
            if col.get('hidden') == '1':
                hidden.add(i)
    return widths, hidden, default_w


def col_px(width):
    # Lebar karakter Excel -> piksel (font default Calibri/Arial, max digit width 7 px).
    return int(width * 7 + 5) if width > 0 else 0


def pt_px(pt):
    return round(pt * 96 / 72, 2)


def side(s, palette):
    if s is None or not s.style:
        return None
    return f"{BORDER_CSS.get(s.style, '1px solid')} {color_of(s.color, palette, '#000000')}"


def main(xlsx):
    os.makedirs(OUT_PUBLIC, exist_ok=True)
    shutil.copy(xlsx, os.path.join(OUT_PUBLIC, 'floq-template.xlsx'))
    palette = theme_palette(xlsx)
    wb = openpyxl.load_workbook(xlsx)
    with zipfile.ZipFile(xlsx) as z:
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        wbx = ET.fromstring(z.read('xl/workbook.xml'))
    rel_target = {r.get('Id'): r.get('Target') for r in rels}
    sheet_path = {}
    for s in wbx.find('m:sheets', NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        sheet_path[s.get('name')] = 'xl/' + rel_target[rid]

    out = {}
    for ws in wb.worksheets:
        key, area = SHEETS[ws.title]
        c1, r1, c2, r2 = range_boundaries(area)
        widths, hidden_cols, default_w = raw_cols(xlsx, sheet_path[ws.title])
        cols = [0 if c in hidden_cols else col_px(widths.get(c, default_w)) for c in range(c1, c2 + 1)]
        default_h = ws.sheet_format.defaultRowHeight or 15
        rows = []
        for r in range(r1, r2 + 1):
            d = ws.row_dimensions.get(r)
            rows.append(0 if (d is not None and d.hidden) else pt_px(d.height if d is not None and d.height else default_h))

        merges = {}
        covered = set()
        for mr in ws.merged_cells.ranges:
            mc1, mr1, mc2, mr2 = mr.bounds
            merges[(mr1, mc1)] = (mr2 - mr1 + 1, mc2 - mc1 + 1, mr1, mc1, mr2, mc2)
            for rr in range(mr1, mr2 + 1):
                for cc in range(mc1, mc2 + 1):
                    if (rr, cc) != (mr1, mc1):
                        covered.add((rr, cc))

        styles, style_idx, cells = [], {}, {}

        def style_for(cell, border):
            f, al = cell.font, cell.alignment
            fill = None
            if cell.fill is not None and cell.fill.fill_type == 'solid':
                fill = color_of(cell.fill.fgColor, palette)
            st = {
                'ff': f.name or 'Calibri',
                'sz': f.sz or 11,
                'b': bool(f.b),
                'i': bool(f.i),
                'u': bool(f.u),
                'c': color_of(f.color, palette, '#000000'),
                'bg': fill,
                'h': al.horizontal or 'general',
                'v': al.vertical or 'bottom',
                'w': bool(al.wrap_text),
                'ind': al.indent or 0,
                'rot': al.text_rotation or 0,
                'fmt': cell.number_format or 'General',
                **border,
            }
            k = json.dumps(st, sort_keys=True)
            if k not in style_idx:
                style_idx[k] = len(styles)
                styles.append(st)
            return style_idx[k]

        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                if (r, c) in covered:
                    continue
                cell = ws.cell(r, c)
                m = merges.get((r, c))
                if m:
                    _, _, mr1, mc1, mr2, mc2 = m
                    bt = next((side(ws.cell(mr1, cc).border.top, palette) for cc in range(mc1, mc2 + 1) if side(ws.cell(mr1, cc).border.top, palette)), None)
                    bb = next((side(ws.cell(mr2, cc).border.bottom, palette) for cc in range(mc1, mc2 + 1) if side(ws.cell(mr2, cc).border.bottom, palette)), None)
                    bl = next((side(ws.cell(rr, mc1).border.left, palette) for rr in range(mr1, mr2 + 1) if side(ws.cell(rr, mc1).border.left, palette)), None)
                    br = next((side(ws.cell(rr, mc2).border.right, palette) for rr in range(mr1, mr2 + 1) if side(ws.cell(rr, mc2).border.right, palette)), None)
                else:
                    b = cell.border
                    bt, bb, bl, br = side(b.top, palette), side(b.bottom, palette), side(b.left, palette), side(b.right, palette)
                border = {k: v for k, v in (('bt', bt), ('bb', bb), ('bl', bl), ('br', br)) if v}
                has_style = cell.has_style and (cell.value is not None or border or (cell.fill is not None and cell.fill.fill_type == 'solid'))
                value = cell.value
                if value is None and not has_style and not m:
                    continue
                entry = {}
                if has_style or m or value is not None:
                    entry['s'] = style_for(cell, border)
                if m:
                    entry['span'] = [m[0], m[1]]
                if isinstance(value, str) and value.startswith('='):
                    entry['f'] = value[1:]
                elif value is not None:
                    entry['v'] = value if isinstance(value, (int, float)) else str(value)
                cells[f'{get_column_letter(c)}{r}'] = entry

        images = []
        for img in ws._images:
            a = img.anchor
            fr = a._from

            def pos(col, colOff, row, rowOff):
                x = sum(cols[: max(0, col + 1 - c1)]) + colOff / 9525
                y = sum(rows[: max(0, row + 1 - r1)]) + rowOff / 9525
                return x, y

            x, y = pos(fr.col, fr.colOff, fr.row, fr.rowOff)
            to = getattr(a, 'to', None)
            if to is not None:
                x2, y2 = pos(to.col, to.colOff, to.row, to.rowOff)
            else:
                x2, y2 = x + a.ext.width / 9525, y + a.ext.height / 9525
            ext = img.format.lower() if getattr(img, 'format', None) else 'png'
            name = f'{key}-{len(images) + 1}.{"jpg" if ext in ("jpeg", "jpg") else ext}'
            with open(os.path.join(OUT_PUBLIC, name), 'wb') as fh:
                fh.write(img._data())
            images.append({'src': f'/templates/{name}', 'x': round(x, 1), 'y': round(y, 1), 'w': round(x2 - x, 1), 'h': round(y2 - y, 1)})

        cf = []
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for rule in rules:
                font = rule.dxf.font if rule.dxf is not None else None
                cf.append({'ref': str(rng.sqref), 'op': rule.operator, 'formula': list(rule.formula), 'color': color_of(font.color, palette) if font is not None and font.color is not None else None})

        out[key] = {
            'name': ws.title,
            'area': {'c1': c1, 'r1': r1, 'c2': c2, 'r2': r2},
            'cols': cols,
            'rows': rows,
            'cells': cells,
            'styles': styles,
            'images': images,
            'cf': cf,
            'landscape': ws.page_setup.orientation == 'landscape',
        }

    with open(OUT_JSON, 'w') as fh:
        json.dump(out, fh, separators=(',', ':'), ensure_ascii=False)
    print('ok', {k: (len(v['cells']), len(v['styles']), len(v['images'])) for k, v in out.items()})


if __name__ == '__main__':
    main(sys.argv[1])
