#!/usr/bin/env python3
"""
Tata letak sheet template Laporan Operasional SPBU (public/templates/ops-template.xlsx) untuk PDF yang
identik dengan Excel format asli: lebar kolom, tinggi baris, sel gabungan, garis, isi warna, dan font.
Memakai fungsi pembaca yang sama dengan scripts/build-report-templates.py.

    python3 scripts/build-ops-layouts.py

Hasil: src/data/ops-layouts.json (dimuat hanya saat PDF dibuat).
"""
import importlib.util
import json
import os
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
spec = importlib.util.spec_from_file_location('brt', os.path.join(ROOT, 'scripts', 'build-report-templates.py'))
brt = importlib.util.module_from_spec(spec)
spec.loader.exec_module(brt)

XLSX = os.path.join(ROOT, 'public', 'templates', 'ops-template.xlsx')
OUT_JSON = os.path.join(ROOT, 'src', 'data', 'ops-layouts.json')
# Sheet yang diunduh dari modulnya masing-masing (tidak ikut laporan operasional).
LEWATI = {'Cek Terra Harian', 'Catatan Kualitas Harian', 'Catatan Persediaan BBM'}


def layout(ws, xlsx, sheet_path, palette):
    c1, r1, c2, r2 = range_boundaries(ws.dimensions)
    c1, r1 = 1, 1
    widths, hidden_cols, default_w = brt.raw_cols(xlsx, sheet_path)
    cols = [0 if c in hidden_cols else brt.col_px(widths.get(c, default_w)) for c in range(c1, c2 + 1)]
    default_h = ws.sheet_format.defaultRowHeight or 15
    rows = []
    for r in range(r1, r2 + 1):
        d = ws.row_dimensions.get(r)
        rows.append(0 if (d is not None and d.hidden) else brt.pt_px(d.height if d is not None and d.height else default_h))

    merges, covered = {}, set()
    for mr in ws.merged_cells.ranges:
        mc1, mr1, mc2, mr2 = mr.bounds
        merges[(mr1, mc1)] = (mr2 - mr1 + 1, mc2 - mc1 + 1, mr1, mc1, mr2, mc2)
        for rr in range(mr1, mr2 + 1):
            for cc in range(mc1, mc2 + 1):
                if (rr, cc) != (mr1, mc1):
                    covered.add((rr, cc))

    styles, style_idx, cells = [], {}, {}

    def style_for(cell, border):
        f, al = cell.font, cell.alignment
        fill = brt.color_of(cell.fill.fgColor, palette) if cell.fill is not None and cell.fill.fill_type == 'solid' else None
        st = {
            'ff': f.name or 'Calibri', 'sz': f.sz or 11, 'b': bool(f.b), 'i': bool(f.i), 'u': bool(f.u),
            'c': brt.color_of(f.color, palette, '#000000'), 'bg': fill, 'h': al.horizontal or 'general',
            'v': al.vertical or 'bottom', 'w': bool(al.wrap_text), 'ind': al.indent or 0,
            'rot': al.text_rotation or 0, 'fmt': cell.number_format or 'General', **border,
        }
        k = json.dumps(st, sort_keys=True)
        if k not in style_idx:
            style_idx[k] = len(styles)
            styles.append(st)
        return style_idx[k]

    side = lambda s: brt.side(s, palette)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            m = merges.get((r, c))
            if m:
                _, _, mr1, mc1, mr2, mc2 = m
                bt = next((side(ws.cell(mr1, cc).border.top) for cc in range(mc1, mc2 + 1) if side(ws.cell(mr1, cc).border.top)), None)
                bb = next((side(ws.cell(mr2, cc).border.bottom) for cc in range(mc1, mc2 + 1) if side(ws.cell(mr2, cc).border.bottom)), None)
                bl = next((side(ws.cell(rr, mc1).border.left) for rr in range(mr1, mr2 + 1) if side(ws.cell(rr, mc1).border.left)), None)
                br = next((side(ws.cell(rr, mc2).border.right) for rr in range(mr1, mr2 + 1) if side(ws.cell(rr, mc2).border.right)), None)
            else:
                b = cell.border
                bt, bb, bl, br = side(b.top), side(b.bottom), side(b.left), side(b.right)
            border = {k: v for k, v in (('bt', bt), ('bb', bb), ('bl', bl), ('br', br)) if v}
            solid = cell.fill is not None and cell.fill.fill_type == 'solid'
            has_style = cell.has_style and (cell.value is not None or border or solid)
            value = cell.value
            if value is None and not has_style and not m:
                continue
            entry = {'s': style_for(cell, border)}
            if m:
                entry['span'] = [m[0], m[1]]
            if isinstance(value, str) and value.startswith('='):
                entry['f'] = value[1:]
            elif value is not None:
                entry['v'] = value if isinstance(value, (int, float)) else str(value)
            cells[f'{get_column_letter(c)}{r}'] = entry

    return {
        'name': ws.title,
        'area': {'c1': c1, 'r1': r1, 'c2': c2, 'r2': r2},
        'cols': cols,
        'rows': rows,
        'cells': cells,
        'styles': styles,
        'images': [],
        'cf': [],
        'landscape': True,
    }


def main():
    palette = brt.theme_palette(XLSX)
    wb = openpyxl.load_workbook(XLSX)
    with zipfile.ZipFile(XLSX) as z:
        rels = ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        wbx = ET.fromstring(z.read('xl/workbook.xml'))
    target = {r.get('Id'): r.get('Target') for r in rels}
    path = {}
    for s in wbx.find('m:sheets', brt.NS):
        rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        path[s.get('name')] = 'xl/' + target[rid].lstrip('/').replace('xl/', '')
    out = {ws.title: layout(ws, XLSX, path[ws.title], palette) for ws in wb.worksheets if ws.title not in LEWATI}
    with open(OUT_JSON, 'w') as fh:
        json.dump(out, fh, separators=(',', ':'), ensure_ascii=False)
    print('ok', {k: (len(v['cells']), len(v['rows']), len(v['cols'])) for k, v in out.items()}, os.path.getsize(OUT_JSON))


if __name__ == '__main__':
    main()
